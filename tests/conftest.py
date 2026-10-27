from pathlib import Path

import pytest
from docx import Document

from coredescriptionw2e.smoke import create_docx, create_template


@pytest.fixture
def sample(tmp_path: Path) -> tuple[Path, Path, Path]:
    folder = tmp_path / "Word с пробелами"
    folder.mkdir()
    word = folder / "Учебный.docx"
    template = tmp_path / "Шаблон.xlsx"
    create_docx(word)
    create_template(template)
    return folder, word, template


@pytest.fixture
def multi_sample(sample):
    folder, word, template = sample
    doc = Document()
    doc.add_paragraph("Скв. TEST-MULTI; Месторождение: Учебное")
    for number, top, drilled, recovered, powers in (
        (1, 100, 4, 3, (1, 1, 2)),
        (2, 104, 3, 2.25, (1, 2)),
    ):
        doc.add_paragraph(f"Долбление {number}")
        doc.add_paragraph(f"Инт: {top}-{top + drilled}м по бурению и каротажу")
        doc.add_paragraph(f"Проходка {drilled}м, вынос керна {recovered}м")
        doc.add_paragraph("Пласт Учебный")
        for layer, power in enumerate(powers, 1):
            cell = doc.add_table(rows=1, cols=1).cell(0, 0)
            cell.text = f"Слой {layer}. {top}-{top + power}м ({top}-{top + power}м)"
            cell.add_paragraph(f"{power}м – Песчаник учебный. {power}м от верха.")
            top += power
    doc.save(word)
    return folder, word, template
