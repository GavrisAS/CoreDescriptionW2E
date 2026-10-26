import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from coredescriptionw2e.domain import ConversionError
from coredescriptionw2e.excel.writer import publish, set_text, write_excel
from coredescriptionw2e.parser import parse_document
from coredescriptionw2e.smoke import verify_output
from coredescriptionw2e.word.common import read_word


def test_excel_end_to_end_and_template_unchanged(sample):
    folder, word, template = sample
    before = hashlib.sha256(template.read_bytes()).digest()
    parsed = parse_document(read_word(word))
    output = folder / "output" / "Учебный.xlsx"
    issues = write_excel(parsed, template, word, output)
    verify_output(output)
    assert before == hashlib.sha256(template.read_bytes()).digest()
    assert "template_well_conflict" in [i.code for i in issues]
    assert not list(output.parent.glob(".core-w2e-*"))


@pytest.mark.parametrize(
    "value", ['=HYPERLINK("https://example.invalid")', "+cmd", "-1+2", "@SUM(A1)"]
)
def test_text_injection_stays_literal(sample, value):
    _, _, template = sample
    w = load_workbook(template)
    set_text(w.active["W2"], value)
    w.save(template)
    w.close()
    w = load_workbook(template)
    assert w.active["W2"].value == value
    assert w.active["W2"].data_type == "s"
    assert w.active["W2"].hyperlink is None
    w.close()


def test_reordered_schema_formula_references(sample):
    folder, word, template = sample
    w = load_workbook(template)
    s = w.active
    for row in s:
        row[13].value, row[14].value = row[14].value, row[13].value
    w.save(template)
    w.close()
    output = folder / "reordered.xlsx"
    write_excel(parse_document(read_word(word)), template, word, output)
    w = load_workbook(output)
    assert w.active["M2"].value == "=N2-O2"
    assert w.active["O2"].value == 100
    w.close()


def test_bad_schema_and_truncation_rejected(sample):
    folder, word, template = sample
    parsed = parse_document(read_word(word))
    output = folder / "out.xlsx"
    parsed.layers[0].full_text = "a" * 32768
    with pytest.raises(ConversionError, match="усечение"):
        write_excel(parsed, template, word, output)
    assert not output.exists() and not list(folder.glob(".core-w2e-*"))
    w = load_workbook(template)
    w.active["N1"] = "Неизвестный"
    w.save(template)
    w.close()
    with pytest.raises(ConversionError, match="заголовок"):
        write_excel(parsed, template, word, output)


def test_atomic_no_clobber(tmp_path: Path):
    target, temp = tmp_path / "target", tmp_path / "temp"
    target.write_bytes(b"old")
    temp.write_bytes(b"new")
    with pytest.raises(FileExistsError):
        publish(temp, target, False)
    assert target.read_bytes() == b"old"
    publish(temp, target, True)
    assert target.read_bytes() == b"new" and not temp.exists()


def test_failed_publication_cleans_temp_and_keeps_result(sample, monkeypatch):
    from coredescriptionw2e.excel import writer

    folder, word, template = sample
    output = folder / "result.xlsx"
    output.write_bytes(b"existing")

    def fail_publish(*args):
        raise PermissionError("locked")

    monkeypatch.setattr(writer, "publish", fail_publish)
    with pytest.raises(PermissionError):
        write_excel(parse_document(read_word(word)), template, word, output, overwrite=True)
    assert output.read_bytes() == b"existing"
    assert not list(folder.glob(".core-w2e-*"))
