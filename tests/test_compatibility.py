import hashlib
import json
import subprocess
import sys
import zipfile
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from coredescriptionw2e.excel.writer import schema
from coredescriptionw2e.service import convert_batch


def assert_run_rows(result):
    assert result.status == "success" and result.output is not None and result.parsed is not None
    with zipfile.ZipFile(result.output) as archive:
        assert archive.testzip() is None
    workbook = load_workbook(result.output)
    try:
        sheet = workbook.active
        mapping = schema(sheet)
        assert sheet.max_row == 6
        assert [sheet.cell(row, mapping[19]).value for row in range(2, 7)] == [1, 2, 3, 1, 2]
        row = 2
        for run in result.parsed.runs:
            for layer in run.layers:
                assert sheet.cell(row, mapping[9]).value == (
                    float(run.drilled) if run.drilled is not None else None
                )
                assert sheet.cell(row, mapping[10]).value == (
                    float(run.recovered) if run.recovered is not None else None
                )
                formula = sheet.cell(row, mapping[11])
                if run.drilled is not None and run.recovered is not None:
                    j, k = (get_column_letter(mapping[key]) for key in (9, 10))
                    assert formula.value == f'=IFERROR({k}{row}/{j}{row},"")'
                    assert formula.data_type == "f"
                else:
                    assert formula.value is None
                assert formula.number_format == "0.00%"
                assert sheet.cell(row, mapping[22]).value == layer.full_text
                assert layer.drilling is not None
                for key, value in ((13, layer.drilling.top), (14, layer.drilling.bottom)):
                    assert sheet.cell(row, mapping[key]).value == float(value)
                row += 1
        assert not sheet.merged_cells.ranges
    finally:
        workbook.close()


@pytest.mark.parametrize("reordered", [False, True])
@pytest.mark.parametrize("missing", [None, "drilled", "recovered", "both"])
@pytest.mark.parametrize("missing_run", [0, 1])
def test_all_layers_metrics_mapping_and_missing(multi_sample, reordered, missing, missing_run):
    folder, word, template = multi_sample
    doc = Document(word)
    metric_paragraphs = [p for p in doc.paragraphs if p.text.startswith("Проходка")]
    for index, p in enumerate(metric_paragraphs):
        if index == missing_run:
            drilled, recovered = ((4, 3), (3, 2.25))[index]
            p.text = {
                None: p.text,
                "drilled": f"Вынос керна {recovered}м",
                "recovered": f"Проходка {drilled}м",
                "both": "Показатели не указаны",
            }[missing]
    doc.save(word)
    if reordered:
        workbook = load_workbook(template)
        sheet = workbook.active
        # Move recovery columns as well as layer depths, including the formula column.
        for row in sheet:
            for a, b in ((9, 0), (10, 1), (11, 2), (13, 14)):
                row[a].value, row[b].value = row[b].value, row[a].value
        workbook.save(template)
        workbook.close()
    before = hashlib.sha256(template.read_bytes()).digest()
    source_before = word.read_bytes()
    result = convert_batch(folder, template)[0]
    assert_run_rows(result)
    assert hashlib.sha256(template.read_bytes()).digest() == before
    assert word.read_bytes() == source_before
    assert ("missing_recovery" in {i.code for i in result.issues}) == (missing is not None)
    # Independent expected source values guard against a shared parser/writer mistake.
    assert result.parsed is not None
    assert [len(r.layers) for r in result.parsed.runs] == [3, 2]
    for index, (drilled, recovered) in enumerate(((4, 3), (3, Decimal("2.25")))):
        run = result.parsed.runs[index]
        assert run.drilled == (
            None if index == missing_run and missing in {"drilled", "both"} else drilled
        )
        assert run.recovered == (
            None if index == missing_run and missing in {"recovered", "both"} else recovered
        )
    validation = subprocess.run(
        [
            sys.executable,
            "scripts/local_validation.py",
            "--word",
            str(word),
            "--xlsx",
            str(result.output),
            "--template",
            str(template),
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert validation.returncode == 0, validation.stderr
    assert json.loads(validation.stdout)["checks"]["recovery_all_run_rows"]


def test_numbered_interval_and_paired_recovery_in_docx(multi_sample):
    folder, word, template = multi_sample
    doc = Document(word)
    expected = []
    for table in doc.tables:
        cell = table.cell(0, 0)
        cell.paragraphs[0].text = cell.paragraphs[0].text.replace("Слой", "Интервал")
        # A manual break is a source boundary; preserve both lines in W.
        body = cell.paragraphs[1]
        body.text = body.text.replace("учебный.", "учебный.\nПродолжение описания.")
        expected.append(body.text)
    for p in doc.paragraphs:
        p.text = p.text.replace("вынос керна", "керна")
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert_run_rows(result)
    assert result.parsed is not None
    assert [layer.full_text for layer in result.parsed.layers] == expected
    assert all("tc" in layer.source.position for layer in result.parsed.layers)


@pytest.mark.parametrize("header", ["Керн 7", "Керн № 7"])
def test_unsupported_layout_isolated_and_safe(multi_sample, header):
    folder, _, template = multi_sample
    unsupported = folder / "unsupported.docx"
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    for line in (header, "Инт. 200-204 м", "проходка 4 м", "л.в.к. 2 м", "Описание образца"):
        cell.add_paragraph(line)
    doc.save(unsupported)
    progress = []
    results = convert_batch(
        folder, template, progress=lambda n, total, r: progress.append((n, total))
    )
    bad = next(r for r in results if r.source == unsupported)
    assert bad.status == "error" and bad.output is None
    assert "unsupported_core_layout" in {i.code for i in bad.issues}
    assert not (folder / "output" / "unsupported.xlsx").exists()
    assert len(results) == 2 and progress == [(1, 2), (2, 2)]
    assert_run_rows(next(r for r in results if r.status == "success"))
    assert all(str(folder) not in i.message and header not in i.message for i in bad.issues)


def test_corrupt_docx_isolated_with_safe_code(multi_sample):
    folder, _, template = multi_sample
    corrupt = folder / "broken.docx"
    corrupt.write_bytes(b"invalid ZIP")
    results = convert_batch(folder, template)
    assert len(results) == 2
    bad = next(result for result in results if result.source == corrupt)
    assert bad.status == "error" and bad.output is None
    assert {issue.code for issue in bad.issues} == {"invalid_docx"}
    assert_run_rows(next(result for result in results if result.status == "success"))


@pytest.mark.parametrize("metric", ["Проходка 0м, керна 1м", "Проходка 4м, керна 5м"])
def test_invalid_paired_metrics_do_not_publish(multi_sample, metric):
    folder, word, template = multi_sample
    doc = Document(word)
    next(p for p in doc.paragraphs if p.text.startswith("Проходка")).text = metric
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert result.status == "error" and result.output is None
    assert not (folder / "output").exists()


def test_local_validation_catches_later_row_formula(multi_sample):
    folder, word, template = multi_sample
    result = convert_batch(folder, template)[0]
    assert result.output is not None
    workbook = load_workbook(result.output)
    workbook.active["L3"] = '=IFERROR(K2/J2,"")'
    workbook.save(result.output)
    workbook.close()
    validation = subprocess.run(
        [
            sys.executable,
            "scripts/local_validation.py",
            "--word",
            str(word),
            "--xlsx",
            str(result.output),
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert validation.returncode == 1
    assert not json.loads(validation.stdout)["checks"]["recovery_all_run_rows"]
