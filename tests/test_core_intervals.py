"""All material is invented. No private source or expected source values."""

import json
import struct
import subprocess
import sys
import zipfile
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import load_workbook

from coredescriptionw2e.cli import main
from coredescriptionw2e.domain import ConversionError, ParseIssue
from coredescriptionw2e.excel.writer import schema, write_excel
from coredescriptionw2e.parser import parse_document
from coredescriptionw2e.service import convert_batch
from coredescriptionw2e.smoke import SYNTHETIC, create_core_docx, create_docx, verify_review
from coredescriptionw2e.word.common import make_document, read_word
from coredescriptionw2e.word.doc_reader import extract_main_story
from coredescriptionw2e.word.tables import binary_core_rows
from tests.test_reader import fib_and_piece


@pytest.fixture
def core_sample(sample):
    folder, word, template = sample
    create_core_docx(word)
    return folder, word, template


def test_core_cells_and_control_roundtrip(core_sample):
    folder, word, template = core_sample
    original = word.read_bytes(), template.read_bytes()
    with patch("socket.socket.connect", side_effect=AssertionError("network forbidden")):
        r = convert_batch(folder, template)[0]
    assert r.parsed is not None and r.output is not None
    assert r.status == "review_required" and r.parsed.status == "REVIEW_REQUIRED"
    assert r.parsed.mode == "CORE_INTERVAL" and len(r.parsed.cores) == 3
    assert len(r.parsed.no_core) == 2 and not r.parsed.layers and not r.parsed.runs
    assert [c.number for c in r.parsed.cores] == [1, 2, 2]
    assert r.parsed.cores[0].metrics["linear_recovery"].source.position.endswith("/p:3")
    assert r.output.parent == folder / "output" / "review"
    verify_review(r.output, core=True)
    with zipfile.ZipFile(r.output) as archive:
        assert archive.testzip() is None
        assert not any("externalLinks" in n for n in archive.namelist())
    assert (word.read_bytes(), template.read_bytes()) == original
    book = load_workbook(r.output)
    control = book["Контроль"]
    assert any(c.value == "SOURCE_FIELD" for row in book["Источники"] for c in row)
    assert all(
        not c.hyperlink or c.hyperlink.target.startswith("#") for row in control for c in row
    )
    book.close()


@pytest.mark.parametrize(
    "code",
    [
        "parse_failure",
        "reader_incomplete",
        "invalid_xml",
        "missing_text",
        "ambiguous_boundaries",
        "missing_identifier",
        "unknown_failure",
    ],
)
@pytest.mark.parametrize("severity", ["error", "warning"])
def test_review_cannot_bypass_structural_failure(core_sample, code, severity):
    folder, word, template = core_sample
    parsed = parse_document(read_word(word))
    parsed.issues.append(ParseIssue(code, "Synthetic structural failure", severity))
    if severity == "warning" and code == "unknown_failure":
        # Unknown nonfatal observations remain warnings; unknown errors never pass.
        return
    assert not parsed.exportable
    output = folder / "never.xlsx"
    with pytest.raises(ConversionError):
        write_excel(parsed, template, word, output)
    assert not output.exists()


@pytest.mark.parametrize(
    "change,code",
    [
        ("empty", "missing_text"),
        ("heading", "unclassified_core_row"),
        ("metric", "missing_metric"),
        ("unknown_row", "unclassified_core_row"),
        ("interval", "ambiguous_boundaries"),
    ],
)
def test_incomplete_core_rows_never_export(core_sample, change, code):
    folder, word, template = core_sample
    d = Document(word)
    cells = d.tables[0].rows[0].cells
    if change == "empty":
        cells[1].text = ""
    elif change == "heading":
        cells[0].text = cells[0].text.replace("Керн № 1", "Неизвестный идентификатор")
    elif change == "metric":
        cells[0].text = cells[0].text.replace("л.в.к. 1 м", "неизвестно")
    elif change == "interval":
        cells[0].text += "\nДополнительно 50–60 м"
    else:
        cells = d.tables[0].add_row().cells
        cells[0].text, cells[1].text = "Неизвестные показатели", "Текст не должен теряться"
    d.save(word)
    r = convert_batch(folder, template)[0]
    assert r.output is None and r.status == "unsupported"
    assert code in {i.code for i in r.issues}


@pytest.mark.parametrize("header", ["Керн 1", "КЕРН № 1", "Керн\u00a0№\u00a01"])
def test_core_headings_are_not_layer_numbers(core_sample, header):
    folder, word, template = core_sample
    d = Document(word)
    c = d.tables[0].rows[0].cells[0]
    c.text = c.text.replace("Керн № 1", header)
    d.save(word)
    r = convert_batch(folder, template)[0]
    assert r.parsed is not None
    assert len(r.parsed.cores) == 3 and not r.parsed.layers


def test_percent_threshold_and_recovery_over_100(core_sample):
    _, word, _ = core_sample
    d = Document(word)
    left = d.tables[0].rows[0].cells[0]
    left.text = left.text.replace("% выноса 50", "% выноса 50,5")
    d.save(word)
    p = parse_document(read_word(word))
    assert not any(i.code == "percent_mismatch" and i.source == p.cores[0].source for i in p.issues)
    left.text = left.text.replace("50,5", "50,51").replace("л.в.к. 1 м", "л.в.к. 3 м")
    d.save(word)
    p = parse_document(read_word(word))
    assert p.exportable and p.status == "REVIEW_REQUIRED"
    assert "core_recovery_exceeds_drilled" in {i.code for i in p.issues}


def test_layer_overlap_review_preserves_original_and_blocks_reverse(sample):
    folder, word, template = sample
    d = Document(word)
    for p in d.paragraphs:
        if p.text.startswith("Слой 3"):
            p.text = p.text.replace("102,00-103,00м", "101,75-103,00м", 1)
    d.save(word)
    result = convert_batch(folder, template)[0]
    assert result.status == "review_required"
    assert result.output is not None and result.parsed is not None
    assert result.parsed.layers[1].drilling is not None
    verify_review(result.output, core=False)
    assert result.parsed.layers[1].drilling.top == Decimal("101.75")
    for p in d.paragraphs:
        if p.text.startswith("Слой 3"):
            p.text = p.text.replace("101,75-103,00м", "99-103м", 1)
    d.save(word)
    bad = convert_batch(folder, template, folder / "other")[0]
    assert bad.output is None
    assert "layer_order" in {i.code for i in bad.issues}


def test_declared_core_replaces_boxed_priority_and_zero_conflict():
    lines = list(SYNTHETIC)
    lines[4] = "Проходка 3м, керна 3м, в коробках 2.75 м"
    lines += [
        "Долбление 3",
        "Инт:105–107м по бурению и каротажу",
        "Проходка 2м, керна 2м, в коробках 0",
    ]
    doc = make_document(
        Path("synthetic.docx"), [(str(i), s) for i, s in enumerate(lines)], "\n".join(lines)
    )
    p = parse_document(doc)
    assert not p.exportable  # Changed answer 5/6: contradictory zero/positive recovery blocks.
    assert p.runs[0].recovered == Decimal("3")
    assert p.runs[0].metrics["declared_core"].value == 3
    assert p.runs[0].metrics["boxed_core"].source.position == "4"
    assert "zero_recovery_conflict" in {i.code for i in p.issues}
    lines[4] = "Проходка 3м, в коробках 0"
    doc = make_document(
        Path("synthetic.docx"), [(str(i), s) for i, s in enumerate(lines)], "\n".join(lines)
    )
    assert not parse_document(doc).exportable


def test_core_reordered_columns_long_text_injection_atomic(core_sample):
    folder, word, template = core_sample
    d = Document(word)
    text = "=SUM(A1); " + ("Учебный текст.\n" * 140)
    d.tables[0].rows[0].cells[1].text = text.rstrip("\n")
    d.save(word)
    book = load_workbook(template)
    for row in book.active:
        for a, b in ((9, 0), (10, 1), (11, 2), (22, 3)):
            row[a].value, row[b].value = row[b].value, row[a].value
    book.save(template)
    book.close()
    r = convert_batch(folder, template)[0]
    assert r.output is not None
    book = load_workbook(r.output)
    sheet = book["Лист1"]
    mapping = schema(sheet, output=True)
    assert sheet.cell(2, mapping[22]).value == text.rstrip("\n")
    assert sheet.cell(2, mapping[22]).data_type == "s"
    assert sheet.row_dimensions[2].height == 409
    assert "Описание" in book.sheetnames
    assert sheet.cell(2, mapping[11]).value == '=IFERROR(B2/A2,"")'
    book.close()
    before = r.output.read_bytes()
    assert convert_batch(folder, template)[0].status == "skipped"
    assert r.output.read_bytes() == before


def test_core_cli_review_exit_and_local_validator(core_sample, capsys):
    folder, word, template = core_sample
    assert main(["--input-dir", str(folder), "--template", str(template), "--json"]) == 3
    payload = json.loads(capsys.readouterr().out)
    assert payload["results"][0]["document_status"] == "REVIEW_REQUIRED"
    output = next((folder / "output" / "review").glob("*.xlsx"))
    r = subprocess.run(
        [sys.executable, "scripts/local_validation.py", "--word", str(word), "--xlsx", str(output)],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert r.returncode == 0 and json.loads(r.stdout)["status"] == "REVIEW_REQUIRED"


def test_sequential_pass_review_unsupported(core_sample):
    folder, word, template = core_sample
    create_docx(folder / "pass.docx")
    bad = Document()
    bad.add_table(rows=1, cols=1).cell(0, 0).text = "Керн 1"
    bad.save(folder / "unsupported.docx")
    progress = []
    results = convert_batch(folder, template, progress=lambda n, total, r: progress.append(n))
    assert progress == [1, 2, 3]
    assert {r.status for r in results} == {"success", "review_required", "unsupported"}


def test_synthetic_binary_table_cell_boundaries():
    raw = "Скв. SYN\rКерн № 1\rИнт. 10-12 м\rпроходка 2 м\rл.в.к. 1 м\r% выноса 50\x07Верх: учебный.\rНиз: учебный.\x07\x07Инт. 12-13 м\x07Керн не поднят.\x07\x07"
    word, table = fib_and_piece(raw)
    extracted = extract_main_story(word, table)
    rows = binary_core_rows("synthetic.doc", extracted)
    assert len(rows) == 2
    assert rows[0].cells[0].blocks[0].text == "Керн № 1"
    assert rows[0].cells[1].text == "Верх: учебный.\nНиз: учебный."
    assert rows[0].cells[1].source.position.startswith("cp:")


def textbox_piece(text: str, *, header: bool = False):
    main = "Учебная основная story\r"
    word, table = fib_and_piece(main + text)
    word = bytearray(word)
    struct.pack_into("<I", word, 76, len(main))
    struct.pack_into("<I", word, 64 + (10 if header else 9) * 4, len(text))
    plc = struct.pack("<III", 0, max(1, len(text) - 1), 1000) + bytes(44)
    index = 58 if header else 56
    struct.pack_into("<II", word, 154 + index * 8, len(table), len(plc))
    return bytes(word), table + plc, main


@pytest.mark.parametrize("header", [False, True])
def test_empty_textboxes_have_exhaustive_coverage(header):
    word, table, main = textbox_piece("\r\r", header=header)
    assert extract_main_story(word, table) == main


@pytest.mark.parametrize("text", ["X\r", " \r", "\t\r", "\x08\r", "\x13\r"])
def test_any_textbox_content_stays_blocked(text):
    word, table, _ = textbox_piece(text)
    with pytest.raises(ConversionError) as exc:
        extract_main_story(word, table)
    assert exc.value.code == "unsupported_doc_textbox"


def test_empty_textbox_bad_plcf_and_short_story_blocked():
    word, table, _ = textbox_piece("\r\r")
    with pytest.raises(ConversionError):
        extract_main_story(word, table[:-1])
    malformed = bytearray(table)
    struct.pack_into("<I", malformed, 21 + 4, 5)
    with pytest.raises(ConversionError):
        extract_main_story(word, bytes(malformed))


@pytest.mark.parametrize("binary", [False, True])
def test_layer_side_columns_excluded_but_right_continuations_retained(tmp_path, binary):
    # A service cell must not be appended to the previous W. The rightmost
    # continuation in another physical row must be appended, without loss.
    rows = [
        ["", "Слой 1. 10–12м (10–12м)\n2м – Песчаник учебный."],
        ["=SAMPLE; Учебная подпись", "Продолжение описания. 2м от верха."],
        ["", "Слой 2. 12–13м (12–13м)\n1м – Алевролит учебный. 1м от верха."],
    ]
    preface = ["Долбление 1", "Инт:10–13м по бурению и каротажу", "Проходка 3м, керна 3м"]
    if binary:
        # Byte piece-table fixture, no external office application required.
        raw = "\r".join(preface) + "\r"
        raw += "".join("\x07".join(c.replace("\n", "\r") for c in row) + "\x07\x07" for row in rows)
        # Use the real OLE fixture shell so read_doc also attaches physical cells.
        # A dedicated unit import path exercises its physical-cell helper below.
        from coredescriptionw2e.word.doc_reader import attach_cell_sources

        word, table = fib_and_piece(raw)
        decoded = extract_main_story(word, table)
        import re

        doc = make_document(
            Path("synthetic.doc"),
            [(f"cp:{m.start()}-{m.end()}", m.group()) for m in re.finditer(r"[^\r\x07]+", decoded)],
            decoded,
        )
        doc.blocks = attach_cell_sources(doc.blocks, decoded)
    else:
        source = tmp_path / "source.docx"
        document = Document()
        for line in preface:
            document.add_paragraph(line)
        table = document.add_table(rows=0, cols=2)
        for left, right in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = left, right
        document.save(source)
        doc = read_word(source)
    parsed = parse_document(doc)
    assert parsed.exportable and len(parsed.layers) == 2
    assert (
        parsed.layers[0].full_text == "2м – Песчаник учебный.\nПродолжение описания. 2м от верха."
    )
    assert any(i.code == "ancillary_table_content" for i in parsed.issues)
    assert any("=SAMPLE" in b.raw for b in doc.blocks)


def test_configurable_percentage_threshold(core_sample):
    folder, word, template = core_sample
    doc = Document(word)
    c = doc.tables[0].rows[0].cells[0]
    c.text = c.text.replace("выноса 50", "выноса 50,25")
    doc.save(word)
    parsed = parse_document(read_word(word), percent_tolerance=Decimal("0.1"))
    assert any(
        i.code == "percent_mismatch" and i.source == parsed.cores[0].source for i in parsed.issues
    )
    with pytest.raises(ConversionError):
        parse_document(read_word(word), percent_tolerance=Decimal("-1"))


def test_zero_drilling_is_never_validated_as_no_core():
    lines = [
        *SYNTHETIC,
        "Долбление 3",
        "Инт:105–107м по бурению и каротажу",
        "Проходка 0м, в коробках 0",
    ]
    doc = make_document(
        Path("synthetic.docx"), [(str(i), s) for i, s in enumerate(lines)], "\n".join(lines)
    )
    p = parse_document(doc)
    assert not p.exportable and "invalid_recovery" in {i.code for i in p.issues}


def test_explicit_layered_well_summary_has_own_scope(sample):
    folder, word, template = sample
    d = Document(word)
    for prefix, drilled, recovered in (
        ("Всего по пласту Учебному", "3м", "3м"),
        ("Всего по скважине", "5м", "4.75м"),
    ):
        left, right = d.add_table(rows=1, cols=2).rows[0].cells
        left.text = prefix + " с отбором керна пройдено:\nВсего поднято керна:"
        right.text = drilled + "\n" + recovered
    d.save(word)
    p = parse_document(read_word(word))
    assert p.totals["linear_recovery"].value == Decimal("4.75")
    assert p.status == "REVIEW_REQUIRED" and "summary_mismatch" in {i.code for i in p.issues}
    result = convert_batch(folder, template)[0]
    assert result.output is not None
    book = load_workbook(result.output)
    row = next(row for row in book["Контроль"] if row[11].value == "summary_mismatch")
    assert "Проходка Word" in row[8].value and "По уникальным исходным записям" in row[9].value
    assert "Decimal" not in row[9].value and "{" not in row[9].value
    book.close()


def test_workbook_failure_does_not_report_document_pass(sample):
    from coredescriptionw2e.service import result_summary

    folder, _, template = sample
    book = load_workbook(template)
    book.active["A1"] = "Invalid schema"
    book.save(template)
    book.close()
    r = convert_batch(folder, template)[0]
    assert r.parsed is not None and r.parsed.status == "PASS"
    assert result_summary(r)["document_status"] == "ERROR" and r.output is None


def test_control_fields_are_literal_and_reserved_template_names(core_sample):
    folder, word, template = core_sample
    d = Document(word)
    left = d.tables[0].rows[0].cells[0]
    left.text += '\n=HYPERLINK("https://example.invalid")'
    d.save(word)
    book = load_workbook(template)
    book.active.title = "Контроль"
    book.save(template)
    book.close()
    r = convert_batch(folder, template)[0]
    assert r.output is not None
    book = load_workbook(r.output)
    assert book.active.title == "Контроль" and "Требует проверки" == book.active["A1"].value
    injected = [
        c
        for row in book["Источники"]
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")
    ]
    assert len(injected) >= 1 and injected[0].data_type == "s" and not injected[0].hyperlink
    book.close()
