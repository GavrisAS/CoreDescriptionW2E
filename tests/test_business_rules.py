"""Batch 1 specialist decisions, using exclusively invented geological material."""

import json
import re
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from coredescriptionw2e.domain import SourceRef
from coredescriptionw2e.excel.writer import OUTPUT_HEADER, schema
from coredescriptionw2e.metrics import extract, recovery
from coredescriptionw2e.parser import parse_document
from coredescriptionw2e.semantics import hydrocarbons, lithology
from coredescriptionw2e.service import batch_summary, convert_batch, input_census
from coredescriptionw2e.word.common import make_document, read_word


def core_rows(word, descriptions, *, metrics="Проходка 2 м, керна 1 м", numbers=None):
    doc = Document()
    doc.add_paragraph("Скв. TEST-aР; Месторождение: Учебное")
    table = doc.add_table(rows=0, cols=2)
    for index, description in enumerate(descriptions):
        left, right = table.add_row().cells
        left.text = f"Керн {numbers[index] if numbers else index + 1}\nИнт. 10–12 м\n{metrics}"
        right.text = description
    doc.save(word)


def test_split_complete_parent_parts_x_and_control(sample):
    folder, word, template = sample
    texts = [
        "0,15м – Песчаник учебный.\nПродолжение с линзой аргиллита 0,01 м.\n2,65м – Аргиллит учебный. 2,65м от верха."
    ]
    core_rows(word, texts)
    r = convert_batch(folder, template)[0]
    assert r.parsed is not None and r.output is not None
    core = r.parsed.cores[0]
    assert len(core.parts) == 2
    assert all(p.parent_source == core.source and p.split_proven for p in core.parts)
    assert "\n".join(p.full_text for p in core.parts) == texts[0]
    assert [b for p in core.parts for b in p.body_blocks] == core.body_blocks
    assert re.findall(r"\d+(?:[.,]\d+)?", "\n".join(p.full_text for p in core.parts)) == re.findall(
        r"\d+(?:[.,]\d+)?", texts[0]
    )
    book = load_workbook(r.output)
    main = book["Лист1"]
    assert main.max_column == 24 and main.max_row == 3
    assert main["X1"].value == OUTPUT_HEADER and main["X2"].value == "TEST-aР"
    assert main.auto_filter.ref == "A1:X3" and "$X$3" in str(main.print_area)
    for row, part in enumerate(core.parts, 2):
        assert main.cell(row, 23).value == part.full_text
        assert all(main.cell(row, col).value is None for col in range(13, 21))
        assert main.cell(row, 10).value == 2 and main.cell(row, 11).value == 1
        assert main.cell(row, 12).value == f'=IFERROR(K{row}/J{row},"")'
        assert main.cell(row, 7).hyperlink.target.startswith("#'Реестр'")
    assert r.parsed.reconciliation["exported_unique"]["drilled"] == 2
    assert book["Контроль"]["A1"].value == "Требует проверки"
    assert any(
        "Геологическая приёмка не выполнена" in str(c.value)
        for row in book["Контроль"]
        for c in row
    )
    book.close()


@pytest.mark.parametrize(
    "text,parts,ambiguous",
    [
        ("Верх: песчаник.\nНиз: аргиллит.", 2, False),
        ("Песчаник с тонким прослоем аргиллита, линза на 0,1–0,2 м.", 1, False),
        ("Песчаник.\nАргиллит.", 1, False),
        ("0,5м – Песчаник.", 1, False),
        ("Общее описание.\n0,5м – Песчаник.\n0,5м – Аргиллит.", 1, True),
        ("0,5м – Песчаник.\n0,1м – а.в.о.\n0,5м – Аргиллит.", 1, True),
        ("Верх: песчаник. Низ: аргиллит.", 1, True),
    ],
)
def test_split_proof_and_ambiguity(sample, text, parts, ambiguous):
    _, word, _ = sample
    core_rows(word, [text])
    p = parse_document(read_word(word))
    assert len(p.cores[0].parts) == parts
    assert ("split_ambiguous" in {i.code for i in p.issues}) == ambiguous
    assert "\n".join(part.full_text for part in p.cores[0].parts) == text


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Переслаивание аргиллита и песчаника.", "Аргиллит / Песчаник"),
        ("0,55м – Серпентинит учебный.", "Серпентинит"),
        ("а.в.о.", None),
        ("Песчаник и аргиллит.", None),
        ("Буроватый углистый оттенок, сульфиды.", None),
    ],
)
def test_lithology_conservative_combination(text, expected):
    assert lithology(text) == expected


@pytest.mark.parametrize(
    "text,positive",
    [
        ("Слабый запах нефти.", True),
        ("Выпоты нефти.", True),
        ("Нефтенасыщенный песчаник.", True),
        ("Следы нефти.", True),
        ("Запах нефти отсутствует.", False),
        ("Нефти нет.", False),
        ("Без запаха нефти.", False),
        ("Не нефтенасыщенный.", False),
        ("Ненефтенасыщенный песчаник.", False),
        ("Ненасыщенный нефтью.", False),
        ("Возможен запах нефти.", False),
        ("Нефтенасыщенность сомнительна.", False),
        ("УРД, буроватый оттенок, углистый, сульфиды.", False),
        ("Нефть упоминается в справочнике.", False),
    ],
)
def test_uv_literal_evidence_and_negation(text, positive):
    assert hydrocarbons(text) == (text if positive else None)


@pytest.mark.parametrize(
    "metric,expected,error",
    [
        ("Проходка 2 м, керна 1 м, в коробках 0,9 м", Decimal(1), False),
        ("Проходка 2 м, в коробках 1 м", None, False),
        ("Проходка 2 м, керна 1 м, линейный вынос 1,2 м, в коробках 0,9 м", None, True),
        ("Проходка 2 м, керна -1 м", Decimal(-1), False),
    ],
)
def test_metric_semantic_sources(metric, expected, error):
    m, issues = extract(metric, SourceRef("synthetic.docx", "p:1", 0))
    assert recovery(m) == expected
    assert bool(issues) == error


@pytest.mark.parametrize("declared_j", [None, 2, 0])
def test_zero_recovery_no_fictitious_layer_or_j(sample, declared_j):
    folder, word, template = sample
    core_rows(
        word,
        ["Керн не поднят."],
        metrics=(f"Проходка {declared_j} м, " if declared_j is not None else "") + "в коробках 0 м",
    )
    r = convert_batch(folder, template)[0]
    assert r.parsed is not None
    if declared_j == 0:
        assert not r.parsed.exportable and r.output is None
        return
    assert r.output is not None and not r.parsed.cores and len(r.parsed.no_core) == 1
    book = load_workbook(r.output)
    main = book["Лист1"]
    assert main.max_row == 2 and main["K2"].value == 0 and main["J2"].value == declared_j
    assert main["L2"].value == ('=IFERROR(K2/J2,"")' if declared_j else None)
    assert main["W2"].value is None and main["T2"].value is None
    assert main["K2"].fill.fgColor.rgb == "00DDEBF7"
    book.close()


@pytest.mark.parametrize("different", [None, "description", "metrics", "context", "interval"])
def test_full_duplicate_requires_all_fields(sample, different):
    folder, word, template = sample
    core_rows(word, ["Песчаник учебный.", "Песчаник учебный."], numbers=[1, 8])
    doc = Document(word)
    row = doc.tables[0].rows[1]
    if different == "description":
        row.cells[1].text += " Нефти нет."
    if different == "metrics":
        row.cells[0].text = row.cells[0].text.replace("керна 1", "керна 0,8")
    if different == "context":
        row.cells[0].text += "\nОтдельное исходное примечание."
    if different == "interval":
        row.cells[0].text = row.cells[0].text.replace("10–12", "12–14")
    doc.save(word)
    r = convert_batch(folder, template)[0]
    assert r.parsed is not None and r.output is not None
    assert len(r.parsed.cores) == 2
    assert len(r.parsed.suppressed_duplicates) == (0 if different else 1)
    assert len(r.parsed.active_cores) == (2 if different else 1)
    if not different:
        assert r.parsed.suppressed_duplicates[0].source_number == 8
        assert r.parsed.reconciliation["exported_unique"]["drilled"] == 2
        book = load_workbook(r.output)
        rows = [row for row in book["Контроль"] if row[11].value == "suppressed_duplicate"]
        assert len(rows) == 1 and rows[0][7].hyperlink
        book.close()


def test_seven_candidates_every_outcome_and_private_audit(sample):
    folder, word, template = sample
    shutil.copyfile(word, folder / "Точки. и РЕГИСТР.DOCX")
    shutil.copyfile(Path("tests/fixtures/synthetic.doc"), folder / "binary.DOC")
    shutil.copyfile(word, folder / "collision.docx")
    shutil.copyfile(Path("tests/fixtures/synthetic.doc"), folder / "collision.doc")
    unknown = Document()
    unknown.add_paragraph("Керн 1")
    unknown.save(folder / "unsupported.docx")
    (folder / "~$locked.docx").write_bytes(b"temporary")
    progress = []
    results = convert_batch(
        folder, template, progress=lambda n, total, r: progress.append((n, total))
    )
    assert len(input_census(folder)) == 7 and len(results) == 7
    assert progress == [(n, 7) for n in range(1, 8)]
    assert batch_summary(results) == {
        "discovered": 7,
        "queued": 6,
        "converted_pass": 3,
        "converted_review": 0,
        "unsupported": 1,
        "error": 2,
        "skipped_with_reason": 1,
    }
    report = json.loads(next((folder / "output").glob("batch-*.json")).read_text())
    assert len(report["census"]) == len(report["results"]) == 7
    assert len(report["private_audit"]) == 7
    assert all(r.output or r.issues for r in results)


def test_control_links_both_conflicting_layers_and_pass_russian(sample):
    folder, word, template = sample
    initial = convert_batch(folder, template)[0]
    assert initial.output is not None
    book = load_workbook(initial.output)
    assert book["Контроль"]["A1"].value == "Проверено автоматически"
    book.close()

    doc = Document(word)
    for paragraph in doc.paragraphs:
        paragraph.text = paragraph.text.replace("102,00-103,00м", "101,75-103,00м")
    doc.save(word)
    result = convert_batch(folder, template, folder / "conflicts")[0]
    assert result.output is not None
    book = load_workbook(result.output)
    assert (
        book["Лист1"]["N2"].fill.fgColor.rgb == book["Лист1"]["N3"].fill.fgColor.rgb == "00FFF2CC"
    )
    issue = next(row for row in book["Контроль"] if row[11].value == "layer_overlap")
    assert "слой 1" in issue[4].value and "слой 3" in issue[4].value
    assert issue[7].hyperlink.target.endswith("W2") and issue[13].hyperlink.target.endswith("W3")
    assert book["Лист1"]["G2"].hyperlink.target == book["Лист1"]["G3"].hyperlink.target
    schema(book["Лист1"], output=True)
    book.close()


def test_cross_stratum_keeps_one_layer_and_all_text():
    lines = [
        "Долбление 1",
        "Инт:10–12м по бурению и каротажу",
        "Проходка 2м, керна 2м",
        "Пласт Учебный-А",
        "Слой 1. 10–12м (10–12м)",
        "2м – Песчаник учебный.",
        "Пласт Учебный-Б",
        "Продолжение того же слоя. 2м от верха.",
    ]
    parsed = parse_document(
        make_document(
            Path("synthetic.docx"),
            [(str(i), text) for i, text in enumerate(lines)],
            "\n".join(lines),
        )
    )
    assert len(parsed.layers) == 1 and parsed.layers[0].stratum is None
    assert parsed.layers[0].full_text == "\n".join(lines[5:])
    assert parsed.status == "REVIEW_REQUIRED"
    assert "stratum_ambiguous" in {i.code for i in parsed.issues}


def test_template_remains_23_and_output_requires_explicit_24_validation(sample):
    folder, _, template = sample
    before = template.read_bytes()
    result = convert_batch(folder, template)[0]
    assert template.read_bytes() == before and result.output is not None
    book = load_workbook(result.output)
    from coredescriptionw2e.domain import ConversionError

    with pytest.raises(ConversionError):
        schema(book["Лист1"])
    assert len(schema(book["Лист1"], output=True)) == 23
    book.close()


def test_independent_private_acceptance_includes_suppressed_source(sample, monkeypatch):
    import sys

    from scripts.business_acceptance import main

    folder, word, template = sample
    core_rows(word, ["0,5м – Песчаник.\n0,5м – Аргиллит."] * 2)
    output = folder / "acceptance"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "business_acceptance",
            "--input-dir",
            str(folder),
            "--template",
            str(template),
            "--output-dir",
            str(output),
        ],
    )
    assert main() == 0
    report = json.loads((output / "report.json").read_text())
    doc = report["documents"][0]
    assert doc["coverage_status"] == "PASS"
    assert doc["independent_coverage"]["source_descriptions"] == 2
    assert doc["coverage"]["registered_duplicate_numeric_tokens"] == 2
    assert len(doc["suppressed_duplicates"]) == 1


@pytest.mark.parametrize("metrics", ["Проходка -1 м, керна 1 м", "Проходка 2 м, керна -1 м"])
def test_negative_metrics_never_publish(sample, metrics):
    folder, word, template = sample
    core_rows(word, ["Песчаник учебный."], metrics=metrics)
    result = convert_batch(folder, template)[0]
    assert result.output is None and result.parsed is not None and not result.parsed.exportable


def test_well_is_not_inferred_from_filename(sample):
    folder, word, template = sample
    core_rows(word, ["Песчаник учебный."])
    doc = Document(word)
    doc.paragraphs[0].text = "Описание учебного керна без идентификатора"
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert result.output is not None
    book = load_workbook(result.output)
    assert book["Лист1"]["D2"].value is None and book["Лист1"]["X2"].value is None
    book.close()


def test_negative_depth_sign_is_never_silently_discarded(sample):
    folder, word, template = sample
    core_rows(word, ["Песчаник учебный."])
    doc = Document(word)
    cell = doc.tables[0].rows[0].cells[0]
    cell.text = cell.text.replace("10–12", "-10–12")
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert result.output is None and result.parsed is not None
    assert result.parsed.cores[0].drilling.top == -10
    assert "invalid_interval" in {i.code for i in result.issues}


@pytest.mark.parametrize("drilled", [None, 2])
def test_layered_zero_only_book_is_not_a_fictitious_layer(sample, drilled):
    folder, word, template = sample
    doc = Document()
    for line in [
        "Скв. SYN-ZERO",
        "Долбление 1",
        "Инт.10–12м по бурению и каротажу",
        (f"Проходка {drilled}м, " if drilled else "") + "в коробках 0",
    ]:
        doc.add_paragraph(line)
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert result.output is not None and result.parsed is not None and not result.parsed.layers
    book = load_workbook(result.output)
    assert book["Лист1"].max_row == 2 and book["Лист1"]["K2"].value == 0
    assert book["Лист1"]["J2"].value == drilled and book["Лист1"]["W2"].value is None
    book.close()


def test_no_core_evidence_cannot_hide_geological_text(sample):
    folder, word, template = sample
    core_rows(word, ["Песчаник учебный."])
    doc = Document(word)
    cells = doc.tables[0].add_row().cells
    cells[0].text = "Инт.12–14 м"
    cells[1].text = "Керн не поднят. Песчаник с содержательным описанием."
    doc.save(word)
    result = convert_batch(folder, template)[0]
    assert result.output is None and result.status == "unsupported"
    assert "ambiguous_boundaries" in {i.code for i in result.issues}


def test_uncertain_uv_reaches_review_and_empty_v(sample):
    folder, word, template = sample
    core_rows(word, ["Песчаник. Нефтенасыщенность сомнительна."])
    result = convert_batch(folder, template)[0]
    assert result.output is not None
    assert "ambiguous_hydrocarbons" in {i.code for i in result.issues}
    book = load_workbook(result.output)
    assert book["Лист1"]["V2"].value is None
    assert "сомнительна" in book["Лист1"]["W2"].value
    assert any(row[11].value == "ambiguous_hydrocarbons" for row in book["Контроль"])
    book.close()
