"""Russian review surface, source registry and reversible internal navigation."""

import re
import textwrap
from collections import Counter, defaultdict
from datetime import UTC, datetime

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from coredescriptionw2e import __version__
from coredescriptionw2e.domain import CoreInterval, CorePart, Layer, NoCoreInterval
from coredescriptionw2e.validation import REVIEW_CODES

STATUS_RU = {
    "PASS": "Проверено автоматически",
    "REVIEW_REQUIRED": "Требует проверки",
    "UNSUPPORTED": "Не удалось обработать",
    "ERROR": "Не удалось обработать",
}
EXPECTED = {
    "missing_stratum",
    "missing_field",
    "unconfirmed_metadata",
    "ambiguous_lithology",
    "missing_end_marker",
    "missing_thickness",
    "ancillary_table_content",
}


def append_safe(sheet, values) -> int:
    from .writer import set_text

    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for col, value in enumerate(values, 1):
        cell = sheet.cell(row, col)
        if isinstance(value, str):
            set_text(cell, value)
        else:
            cell.value = float(value) if hasattr(value, "as_tuple") else value
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[row].height = 48
    return row


def internal_link(cell, title: str, address: str) -> None:
    cell.hyperlink = f"#'{title.replace(chr(39), chr(39) * 2)}'!{address}"
    cell.font = Font(color="0563C1", underline="single")


def interval_text(interval) -> str:
    return f"{interval.top}–{interval.bottom} м" if interval else "Интервал не указан"


def physical_locator(source) -> str:
    if source is None:
        return "Заголовок/итоги документа"
    table = re.search(r"(?:w:tbl|tbl)[\[:](\d+)", source.position)
    row = re.search(r"(?:w:tr|tr)[\[:](\d+)", source.position)
    cell = re.search(r"(?:w:tc|tc)[\[:](\d+)", source.position)
    if (table or "w:tbl" in source.position) and (row or "w:tr" in source.position):
        number = cell[1] if cell else None
        if number is None and (m := re.search(r"/cell:(\d+)", source.position)):
            number = str(int(m[1]) + 1)
        return f"Таблица {table[1] if table else '1'}, строка {row[1] if row else '1'}" + (
            f", ячейка {number}" if number else ""
        )
    return f"Текстовый блок {source.order + 1}"


def locator(owner, entry=None) -> str:
    if isinstance(owner, NoCoreInterval):
        return (
            ("Керн не поднят" if owner.kind == "NO_RECOVERY" else "Бурение без отбора керна")
            + ", "
            + interval_text(owner.drilling)
        )
    if isinstance(owner, CoreInterval):
        return f"Керн {owner.number}, {interval_text(owner.drilling)}" + (
            f", часть {entry.ordinal}" if isinstance(entry, CorePart) else ""
        )
    return f"Долбление {owner.number}" + (
        f", слой {entry.number}, {interval_text(entry.drilling)}"
        if isinstance(entry, Layer)
        else f", {interval_text(owner.drilling)}"
    )


def issue_values(issue, owners, parsed):
    """Original and comparison facts kept separate; never a substituted metric."""
    if issue.code.startswith("summary"):
        source = (
            "; ".join(
                f"{'Проходка Word' if key == 'declared_drilled' else 'Поднято / вынос Word'}: {metric.value} м"
                for key, metric in parsed.totals.items()
            )
            or "Итог Word не определён"
        )
        totals = parsed.reconciliation.get(
            "exported_unique", parsed.reconciliation.get("all_runs", {})
        )
        drilled = totals.get("drilled", totals.get("declared_drilled"))
        recovered = totals.get("recovered", totals.get("linear_recovery"))
        compared = f"По уникальным исходным записям: J={drilled if drilled is not None else 'не определено'} м; K={recovered if recovered is not None else 'не определено'} м. Объём и сопоставимость — проверить по Реестру."
        return source, compared
    if not owners:
        return "См. исходный блок на листе Источники", None
    owner = owners[0][0]
    entry = owners[0][1]
    if issue.code in {"ambiguous_lithology", "ambiguous_hydrocarbons", "split_ambiguous"}:
        return (
            entry.full_text[:240] if entry else "См. исходный блок"
        ), "Неоднозначный признак не подставлен; полное описание сохранено в W."
    if issue.code == "missing_stratum":
        return "Однозначный пласт не определён по описанию", "S оставлена пустой"
    if issue.code in {"layer_overlap", "core_overlap"}:
        return " / ".join(
            interval_text(entry.drilling if isinstance(entry, Layer) else parent.drilling)
            for parent, entry in owners
        ), "Перекрытие; исходные границы не исправлены"
    if "percent" in issue.code:
        return str(
            owner.metrics.get("declared_percent").value
        ) if "declared_percent" in owner.metrics else None, str(
            owner.metrics.get("computed_percent").value
        ) if "computed_percent" in owner.metrics else None
    names = {
        "declared_drilled": "Проходка Word",
        "declared_core": "Керна Word",
        "linear_recovery": "Линейный вынос Word",
        "boxed_core": "В коробках Word",
        "declared_percent": "Процент Word",
    }
    return (
        "; ".join(
            f"{names.get(m.variant, m.variant)}: {m.value}"
            for m in owner.metrics.values()
            if m.status == "DECLARED"
        ),
        f"J={owner.drilled if owner.drilled is not None else 'не указано'}; K={owner.recovered if owner.recovered is not None else 'не указано'}",
    )


def add_control(workbook, parsed, main, mapping, entries, long_texts, writer_issues) -> None:
    control = workbook.create_sheet("Контроль")
    registry = workbook.create_sheet("Реестр")
    sources = workbook.create_sheet("Источники")
    review = parsed.status == "REVIEW_REQUIRED"
    append_safe(control, [STATUS_RU[parsed.status]])
    control["A1"].font = Font(bold=True, color="9C0006" if review else "006100", size=16)
    control.merge_cells("A1:E1")
    control.merge_cells("F1:G1")
    control["F1"] = parsed.status
    control.merge_cells("H1:N1")
    control["H1"] = "CoreDescriptionW2E — контроль переноса"
    append_safe(
        control,
        [
            "Word",
            parsed.document.file_id,
            None,
            None,
            None,
            "Режим",
            parsed.mode,
            None,
            None,
            "Версия",
            __version__,
            "UTC",
            datetime.now(UTC).isoformat(timespec="seconds"),
        ],
    )
    control.merge_cells("B2:E2")
    control.merge_cells("G2:I2")
    control.merge_cells("M2:N2")
    for text in (
        "Геологическая приёмка не выполнена. Визуальная приёмка: PENDING. Исходные значения не исправлялись.",
        "Статусы: Проверено автоматически — структурные проверки пройдены; Требует проверки — исходные конфликты/ограничения; Не удалось обработать — книга не выдаётся.",
        "Цвета: жёлтый — конфликт интервалов; сиреневый — исключённый дубль; голубой — нулевой вынос; оранжевый — отсутствующие данные; красный — расхождение итогов.",
        "J — объявленная проходка; K — керна/линейный вынос (коробки отдельно); L — формула K/J. J/K повторяются по слоям: НЕ СУММИРОВАТЬ строки, итоги — по родителям в Реестре.",
        "Пустые M/N–Q/T частей керна — нет доказанных абсолютных границ/номера слоя. R без независимой ГИС-привязки пусто. Порядок части не является номером слоя.",
    ):
        row = append_safe(control, [text])
        control.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
        control.row_dimensions[row].height = 32
    counts_row = append_safe(
        control,
        [
            f"Родительских кернов: {len(parsed.cores)}; экспортируемых частей: {sum(len(c.parts) for c in parsed.active_cores)}; неразделённых неоднозначных: {sum(i.code == 'split_ambiguous' for i in parsed.issues)}; исключённых дублей: {len(parsed.suppressed_duplicates)}; строк результата: {len(entries)}. Схема core24-v1."
        ],
    )
    control.merge_cells(start_row=counts_row, start_column=1, end_row=counts_row, end_column=14)
    header_row = append_safe(
        control,
        [
            "№",
            "Уровень",
            "Проблема",
            "Word-файл",
            "Источник (найти в Word)",
            "Интервал",
            "Строки Excel",
            "К описанию",
            "Исходное значение",
            "Вычисленное / сравнение",
            "Действие для проверки",
            "issue_code",
            "SourceRef",
            "Другие строки конфликта",
        ],
    )
    append_safe(
        registry,
        ["Реестр исходных записей — суммы только по уникальным родителям; части не складывать"],
    )
    append_safe(
        registry,
        [
            "Источник",
            "Интервал отбора",
            "Строки Excel",
            "J объявлено",
            "K принято",
            "% Word",
            "% K/J",
            "Решение",
            "SourceRef",
            "Источник J",
            "Источник K",
            "Исходные коробки",
            "Частей",
        ],
    )
    append_safe(
        sources,
        [
            "Источник / родитель",
            "Вид",
            "Порядок части",
            "Начало",
            "Конец",
            "SourceRef",
            "Исходный текст / фрагмент (без усечения)",
        ],
    )
    refs = defaultdict(list)
    rows_by_owner = defaultdict(list)
    destination = {
        row: f"{get_column_letter(mapping[22 if entry else 10])}{row}"
        for row, (_, entry) in enumerate(entries, 2)
    }
    for row, (owner, entry) in enumerate(entries, 2):
        rows_by_owner[owner.source].append(row)
        for ref in {
            owner.source,
            entry.source if entry else owner.source,
            *(m.source for m in owner.metrics.values()),
        }:
            refs[ref].append((row, owner, entry))
        if entry:
            for block in entry.body_blocks:
                if not any(item[0] == row for item in refs[block.source]):
                    refs[block.source].append((row, owner, entry))
    owners = [*parsed.runs, *parsed.cores, *parsed.no_core]
    for duplicate in parsed.suppressed_duplicates:
        refs[duplicate.source] = refs[duplicate.retained_source]
    owners.sort(key=lambda item: item.source.order)
    owner_registry = {}
    for owner in owners:
        rows = rows_by_owner[owner.source]
        duplicate = next(
            (d for d in parsed.suppressed_duplicates if d.source == owner.source), None
        )
        retained_rows = rows_by_owner[duplicate.retained_source] if duplicate else rows
        decision = (
            "Исключено как полный дубль"
            if duplicate
            else "Нулевой вынос; геологический слой не создан"
            if (isinstance(owner, NoCoreInterval) and owner.kind == "NO_RECOVERY")
            or getattr(owner, "no_core", False)
            else "Бурение без отбора керна; только Контроль"
            if isinstance(owner, NoCoreInterval)
            else "Учтено один раз в итогах"
        )
        metrics = owner.metrics
        k_source = next(
            (metrics[k] for k in ("linear_recovery", "declared_core") if k in metrics), None
        )
        r = append_safe(
            registry,
            [
                locator(owner),
                interval_text(owner.drilling),
                ", ".join(map(str, retained_rows)),
                owner.drilled,
                owner.recovered,
                metrics["declared_percent"].value if "declared_percent" in metrics else None,
                metrics["computed_percent"].value if "computed_percent" in metrics else None,
                decision,
                owner.source.position,
                metrics["declared_drilled"].source.position
                if "declared_drilled" in metrics
                else None,
                k_source.source.position
                if k_source
                else "Явное свидетельство нуля"
                if owner.recovered == 0
                else None,
                metrics["boxed_core"].value if "boxed_core" in metrics else None,
                len(owner.parts)
                if isinstance(owner, CoreInterval)
                else len(getattr(owner, "layers", [])),
            ],
        )
        owner_registry[owner.source] = r
        if retained_rows:
            internal_link(
                registry.cell(r, 3),
                main.title,
                destination[retained_rows[0]],
            )
        for main_row in rows:
            internal_link(main.cell(main_row, mapping[6]), registry.title, f"A{r}")
        if duplicate:
            for c in registry[r]:
                c.fill = PatternFill("solid", fgColor="E4DFEC")
        if isinstance(owner, CoreInterval):
            for block in [*owner.source_fields, *owner.body_blocks]:
                append_safe(
                    sources,
                    [
                        locator(owner),
                        "Исходный блок",
                        None,
                        None,
                        None,
                        block.source.position,
                        block.raw,
                    ],
                )
            for part in owner.parts:
                append_safe(
                    sources,
                    [
                        locator(owner),
                        "Карта фрагментов W",
                        part.ordinal,
                        part.start,
                        part.end,
                        part.source.position,
                        part.full_text,
                    ],
                )
        elif isinstance(owner, NoCoreInterval):
            append_safe(
                sources,
                [
                    locator(owner),
                    owner.kind,
                    None,
                    None,
                    None,
                    owner.source.position,
                    owner.evidence,
                ],
            )
    for method, values in parsed.reconciliation.items():
        for key, value in values.items():
            append_safe(registry, ["Сверка по уникальным родителям", method, key, value])
    for duplicate in parsed.suppressed_duplicates:
        append_safe(
            registry,
            [
                "Решение об исключении полного дубля",
                f"Керн {duplicate.source_number} → Керн {duplicate.retained_number}",
                ", ".join(duplicate.compared_fields),
                "Изменение J",
                -duplicate.drilled_effect,
                "Изменение K",
                -duplicate.recovered_effect,
                "Оба исходных ID и тексты — на листе Источники",
                duplicate.source.position,
                duplicate.retained_source.position,
            ],
        )
    for key, metric in parsed.totals.items():
        append_safe(
            registry,
            [
                "Итог Word; проверьте сопоставимость семантики",
                key,
                metric.value,
                metric.source.position,
            ],
        )
    for key, note in parsed.reconciliation_notes.items():
        append_safe(registry, ["Сопоставимость итогов", key, note])
    # Preserve all imported source blocks, including service cells and summary.
    for block in parsed.document.blocks:
        append_safe(
            sources,
            ["Исходный Word", "SOURCE_FIELD", None, None, None, block.source.position, block.raw],
        )
    grouped = defaultdict(list)
    for issue in [*parsed.issues, *writer_issues]:
        grouped[
            issue.code
            if issue.code in EXPECTED and issue.severity != "error"
            else (issue.code, issue.source, issue.related_sources)
        ].append(issue)
    ordered_issues = sorted(
        grouped.values(),
        key=lambda items: (
            0
            if "overlap" in items[0].code
            else 1
            if items[0].code in REVIEW_CODES or items[0].severity == "error"
            else 2
        ),
    )
    for number, issues in enumerate(ordered_issues, 1):
        issue = issues[0]
        matches = {}
        for item in issues:
            for ref in (item.source, *item.related_sources):
                for match in refs.get(ref, []):
                    matches[match[0]] = match
        matches = dict(sorted(matches.items()))
        matched_owners = [(owner, entry) for _, owner, entry in matches.values()]
        labels = list(dict.fromkeys(locator(owner, entry) for owner, entry in matched_owners))
        if issue.code == "suppressed_duplicate":
            duplicate = next(d for d in parsed.suppressed_duplicates if d.source == issue.source)
            original_owner = next(c for c in parsed.cores if c.source == duplicate.source)
            labels = [
                f"Исключён: {locator(original_owner)}; оставлен Керн {duplicate.retained_number}"
            ]
        if len(labels) > 3:
            labels = labels[:3] + [
                f"и ещё {len(labels) - 3}; полный перечень — Реестр, строки Excel указаны рядом"
            ]
        source_block = next(
            (
                b
                for b in parsed.document.blocks
                if issue.source and b.source.position == issue.source.position
            ),
            None,
        )
        if not labels:
            owner = next((o for o in owners if o.source == issue.source), None)
            labels = (
                [locator(owner)]
                if owner
                else [
                    physical_locator(issue.source)
                    + ("; ориентир: " + source_block.text[:100] if source_block else "")
                ]
            )
        original, compared = issue_values(issue, matched_owners, parsed)
        r = append_safe(
            control,
            [
                number,
                "Конфликт — проверить"
                if issue.code in REVIEW_CODES or issue.severity == "error"
                else "Предупреждение",
                issue.message + (f" (однотипных: {len(issues)})" if len(issues) > 1 else ""),
                parsed.document.file_id,
                "; ".join(labels),
                " / ".join(
                    dict.fromkeys(
                        interval_text(
                            entry.drilling if isinstance(entry, Layer) else owner.drilling
                        )
                        for owner, entry in matched_owners
                    )
                ),
                ", ".join(map(str, matches)),
                "Открыть описание" if matches else "См. Реестр / Источники",
                original,
                compared,
                "Найти запись по ID/интервалу в Word; сверить значения. Геологическая приёмка вручную.",
                issue.code,
                issue.source.position if issue.source else None,
                ", ".join(map(str, list(matches)[1:])),
            ],
        )
        if matches:
            internal_link(
                control.cell(r, 8),
                main.title,
                destination[next(iter(matches))],
            )
            if len(matches) > 1:
                internal_link(
                    control.cell(r, 14),
                    main.title,
                    destination[list(matches)[1]],
                )
            for main_row in matches:
                if issue.code in REVIEW_CODES:
                    internal_link(main.cell(main_row, mapping[6]), control.title, f"A{r}")
        control.row_dimensions[r].height = 110
        color = (
            "FFF2CC"
            if "overlap" in issue.code
            else "E4DFEC"
            if "duplicate" in issue.code
            else "F8CBAD"
            if "summary" in issue.code
            else "FCE4D6"
        )
        for c in control[r]:
            c.fill = PatternFill("solid", fgColor=color)
    append_safe(control, ["Категории проблем", str(dict(Counter(i.code for i in parsed.issues)))])
    for ws in (control, registry, sources):
        ws.freeze_panes = f"A{header_row + 1}" if ws == control else "C3"
        ws.sheet_view.showGridLines = False
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    for col, width in {
        "A": 6,
        "B": 18,
        "C": 44,
        "D": 24,
        "E": 38,
        "F": 20,
        "G": 12,
        "H": 18,
        "I": 28,
        "J": 28,
        "K": 38,
        "L": 28,
        "M": 28,
        "N": 20,
    }.items():
        control.column_dimensions[col].width = width
    control.column_dimensions["L"].hidden = True
    control.column_dimensions["M"].hidden = True
    registry.merge_cells("A1:M1")
    registry.column_dimensions["A"].width = 52
    registry.column_dimensions["H"].width = 50
    sources.column_dimensions["G"].width = 110
    control.auto_filter.ref = f"A{header_row}:N{control.max_row - 1}"
    control.print_title_rows = f"1:{header_row}"
    for c in control[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="244062")
    if long_texts:
        detail = workbook.create_sheet("Описание")
        append_safe(detail, ["Строка основного листа", "Фрагмент полного W в исходном порядке"])
        detail.column_dimensions["B"].width = 100
        for main_row, text in long_texts:
            for paragraph in text.split("\n"):
                for piece in textwrap.wrap(
                    paragraph, width=85, replace_whitespace=False, drop_whitespace=False
                ) or [""]:
                    r = append_safe(detail, [main_row, piece])
                    internal_link(
                        detail.cell(r, 1), main.title, f"{get_column_letter(mapping[22])}{main_row}"
                    )
    if review:
        workbook.active = workbook.index(control)
