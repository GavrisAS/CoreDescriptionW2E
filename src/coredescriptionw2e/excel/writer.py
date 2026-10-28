"""Versioned 23-column template schema and atomic, formula-safe Excel output."""

import math
import os
import re
import tempfile
import unicodedata
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from coredescriptionw2e.domain import ConversionError, Layer, ParseIssue, ParseResult
from coredescriptionw2e.export_rows import export_rows
from coredescriptionw2e.validation import REVIEW_CODES
from coredescriptionw2e.word.common import MAX_FILE_BYTES, check_file

from .control import add_control

# core23-v1: semantic identifiers, validated by header names, including reordered columns.
HEADERS = (
    "MECTR",
    "PLOSHAD",
    "AREA_CODE",
    "N_SKV",
    "LAS",
    "UWI",
    "ссылка на источник информации",
    "Инт-л долбления Кровля(глуб, м)",
    "Инт-л долбления Подошва(глуб, м)",
    "Проходка, м",
    "линейный Вынос, м",
    "Вынос, %",
    "толщина слоя, м",
    "Кровля слоя по бурению, м",
    "подошва слоя по бурению, м",
    "кровля слоя по ГИС, м",
    "подошва слоя по ГИС, м",
    "Стратиграфическая единица по ГИС",
    "Стратиграфическая единица по керну",
    "Слой-номер",
    "Краткое описание породы",
    "Характер УВ насыщения",
    "Описание слоя",
)


def header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]", "", text)


OUTPUT_HEADER = "Номер скважины по описанию керна"
OUTPUT_SCHEMA = "core24-v1"


def schema(sheet, *, output: bool = False) -> dict[int, int]:
    if sheet.max_column != (24 if output else 23):
        raise ConversionError("Шаблон должен содержать ровно 23 столбца схемы core23-v1.")
    expected = {header_key(value): i for i, value in enumerate(HEADERS)}
    result: dict[int, int] = {}
    if output and sheet.cell(1, 24).value != OUTPUT_HEADER:
        raise ConversionError("Неверный столбец X схемы core24-v1.")
    for cell in sheet[1][:23]:
        key = header_key(cell.value)
        # The legacy template adds a parenthetical explanation to this header.
        if key.startswith(header_key(HEADERS[18])):
            key = header_key(HEADERS[18])
        if key not in expected or expected[key] in result:
            raise ConversionError("Неизвестный или повторяющийся заголовок шаблона core23-v1.")
        result[expected[key]] = cell.column
    if len(result) != 23:
        raise ConversionError("В шаблоне отсутствуют обязательные заголовки.")
    return result


def set_text(cell: Cell, value: str | None) -> None:
    if value is None:
        cell.value = None
        return
    if len(value) > 32767:
        raise ConversionError("Текст превышает лимит ячейки Excel; усечение запрещено.")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise ConversionError("В тексте есть недопустимые для Excel управляющие символы.")
    cell.value = value
    cell.data_type = "s"
    cell.hyperlink = None
    cell.quotePrefix = value.lstrip().startswith(("=", "+", "-", "@"))


def publish(temp: Path, target: Path, overwrite: bool) -> None:
    if overwrite:
        os.replace(temp, target)
    elif os.name == "nt":
        # Windows rename fails if the destination already exists.
        os.rename(temp, target)
    else:
        # replace() would race with a concurrent producer. Atomic no-clobber link.
        os.link(temp, target)
        temp.unlink()


def write_excel(
    parsed: ParseResult, template: Path, source: Path, output: Path, *, overwrite: bool = False
) -> list[ParseIssue]:
    check_file(template)
    if template.suffix.lower() != ".xlsx":
        raise ConversionError("Нужен шаблон XLSX.")
    if output.resolve() in {source.resolve(), template.resolve()}:
        raise ConversionError("Результат совпадает с исходным файлом.")
    if output.is_symlink() or (output.exists() and not overwrite):
        raise ConversionError("Результат уже существует; перезапись не разрешена.")
    if not parsed.exportable:
        raise ConversionError("Критические ошибки разбора; Excel не опубликован.")
    try:
        with zipfile.ZipFile(template) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_FILE_BYTES:
                raise ConversionError("Слишком большой распакованный шаблон.")
        workbook = load_workbook(template, keep_links=False)
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ConversionError("Некорректный XLSX-шаблон.") from exc
    temp: Path | None = None
    try:
        if len(workbook.worksheets) != 1:
            raise ConversionError("Схема core23-v1 поддерживает шаблон с одним листом.")
        sheet = workbook["Лист1"] if "Лист1" in workbook.sheetnames else workbook.active
        if sheet is None:
            raise ConversionError("В шаблоне нет листа.")
        mapping = schema(sheet)
        if sheet.title in {"Контроль", "Описание", "Реестр", "Источники"}:
            sheet.title = "Данные core24-v1"
        issues: list[ParseIssue] = []
        example_wells = {
            str(sheet.cell(r, mapping[3]).value).strip()
            for r in (2, 3)
            if sheet.cell(r, mapping[3]).value is not None
        }
        if parsed.metadata.well and example_wells and example_wells != {parsed.metadata.well}:
            issues.append(
                ParseIssue(
                    "template_well_conflict",
                    "Номер скважины Word отличается от примера Excel; использован только Word.",
                )
            )
        styles = {column: copy(sheet.cell(2, column)._style) for column in mapping.values()}
        for merged in list(sheet.merged_cells.ranges):
            if merged.max_row > 1:
                sheet.unmerge_cells(str(merged))
        sheet.delete_rows(2, max(0, sheet.max_row - 1))
        for row in list(sheet.row_dimensions):
            if row > 1:
                del sheet.row_dimensions[row]
        # Remove stale demonstration table ranges and validation on deleted rows.
        sheet.tables.clear()
        sheet.data_validations.dataValidation.clear()
        sheet.conditional_formatting._cf_rules.clear()
        width = max(
            70, min(100, sheet.column_dimensions[get_column_letter(mapping[22])].width or 70)
        )
        sheet.column_dimensions[get_column_letter(mapping[22])].width = width
        row = 2
        entries = export_rows(parsed)
        affected = {
            ref
            for i in parsed.issues
            if i.code in REVIEW_CODES
            for ref in (i.source, *i.related_sources)
            if ref
        }
        owner_affected = set()
        for owner, _ in entries:
            if any(metric.source in affected for metric in owner.metrics.values()):
                owner_affected.add(owner.source)
        set_text(sheet.cell(1, 24), OUTPUT_HEADER)
        sheet.cell(1, 24)._style = copy(sheet.cell(1, mapping[3])._style)
        sheet.column_dimensions["X"].width = 28
        long_texts = []
        for run, entry in entries:
            for column, style in styles.items():
                sheet.cell(row, column)._style = copy(style)
            values: dict[int, str | int | float | None] = {
                0: parsed.metadata.field_name,
                3: parsed.metadata.well,
                6: os.path.relpath(source.resolve(), output.parent.resolve()),
                9: float(run.drilled) if run.drilled is not None else None,
                10: float(run.recovered) if run.recovered is not None else None,
                20: entry.lithology if entry else None,
                21: entry.hydrocarbons if entry else None,
                22: entry.full_text if entry else None,
            }
            ranges = [(7, run.drilling)]
            if parsed.mode == "LAYERED_DESCRIPTION" and entry is not None:
                # The branch is structural: core IDs never become layer numbers.
                assert isinstance(entry, Layer)
                values.update({18: entry.stratum, 19: entry.number})
                ranges.extend([(13, entry.drilling), (15, entry.gis)])
            set_text(sheet.cell(row, 24), parsed.metadata.well)
            for key, interval in ranges:
                if interval:
                    values[key] = float(interval.top)
                    values[key + 1] = float(interval.bottom)
            for key, value in values.items():
                c = sheet.cell(row, mapping[key])
                if isinstance(value, str):
                    set_text(c, value)
                else:
                    c.value = value

            def address(key: int, row: int = row) -> str:
                return f"{get_column_letter(mapping[key])}{row}"

            if run.drilled is not None and run.drilled > 0 and run.recovered is not None:
                sheet.cell(row, mapping[11], f'=IFERROR({address(10)}/{address(9)},"")')
            if (
                parsed.mode == "LAYERED_DESCRIPTION"
                and isinstance(entry, Layer)
                and entry.drilling is not None
            ):
                sheet.cell(row, mapping[12], f"={address(14)}-{address(13)}")
            for key in range(7, 17):
                sheet.cell(row, mapping[key]).number_format = "0.00%" if key == 11 else "0.00"
            sheet.cell(row, mapping[22]).alignment = Alignment(vertical="top", wrap_text=True)
            # Conservative wrapping estimate, including real newlines; Excel's
            # maximum height is 409 points. W is never changed for presentation.
            lines = sum(
                max(1, math.ceil(len(line.expandtabs(4)) / max(1, int(width * 0.85))))
                for line in (entry.full_text if entry else "").split("\n")
            )
            height = max(30, 16 * lines + 8)
            sheet.row_dimensions[row].height = min(409, height)
            if height > 409:
                long_texts.append((row, entry.full_text if entry else ""))
                issues.append(
                    ParseIssue(
                        "excel_height_limit",
                        "Текст W сохранён целиком; для просмотра см. лист «Описание».",
                        source=entry.source if entry else run.source,
                    )
                )
            if entry is None:
                for c in sheet[row]:
                    c.fill = PatternFill("solid", fgColor="DDEBF7")
            elif entry.source in affected or run.source in affected or run.source in owner_affected:
                for c in sheet[row]:
                    c.fill = PatternFill("solid", fgColor="FFF2CC")
            row += 1
        add_control(workbook, parsed, sheet, mapping, entries, long_texts, issues)
        sheet.auto_filter.ref = f"A1:X{row - 1}"
        sheet.print_area = f"A1:X{row - 1}"
        sheet.freeze_panes = "A2"
        workbook.calculation = CalcProperties(
            calcId=191029, fullCalcOnLoad=True, forceFullCalc=True
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix=".core-w2e-", suffix=".xlsx", dir=output.parent)
        os.close(fd)
        temp = Path(name)
        workbook.save(temp)
        with zipfile.ZipFile(temp) as archive:
            if archive.testzip() is not None:
                raise ConversionError("Не прошла проверка записанного XLSX.")
        with temp.open("r+b") as stream:
            os.fsync(stream.fileno())
        publish(temp, output, overwrite)
        return issues
    finally:
        workbook.close()
        if temp is not None:
            temp.unlink(missing_ok=True)
