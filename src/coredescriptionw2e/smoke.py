"""Synthetic acceptance exercised inside the actual frozen executable."""

import json
import shutil
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import __version__
from .excel.writer import HEADERS
from .service import convert_batch

# Entirely invented material, deliberately unrelated to private input documents.
SYNTHETIC = (
    "Скв. TEST-42; Месторождение: Учебное",
    "Долбление 1",
    "Инт:100,0–103,0м – по бурению",
    "100,5–103,5м – по каротажу",
    "проходка – 3м, вынос керна – 3м",
    "Пласт Учебный (100,0–103,0м)",
    "Слой 1. 100,00—102,00м ((100,50—102,50м))",
    "2,00м – Песчаник серый. Запах",
    "нефти отсутствует. Включение на 0,19–0,23м от начала. 2,00м от верха.",
    "Слой 3. 102,00-103,00м (102,50-103,50м)",
    "=1+1; 1,00м – Переслаивание аргиллита и песчаника. 1,00м от верха.",
    "Всего вынос керна: 3м",
    "Геолог Учебный А.А.",
    "Долбление 2",
    "Инт:103,0-105,0м – по бурению",
    "103,25-105,25м – по каротажу",
    "проходка – 2м, вынос керна – 2м",
    "Доюрский комплекс",
    "Слой 1. 103,00-105,00м (103,25-105,25м)",
    "2,00м – Алевролит серый. 2,00м от верха.",
    "Всего: 5м",
    "Описание керна составил: Учебный А.А.",
)


def create_docx(path: Path) -> None:
    document = Document()
    # Paragraphs, a two-cell table, then ordinary paragraphs: logical order matters.
    for line in SYNTHETIC[:7]:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = SYNTHETIC[7]
    table.cell(0, 1).text = SYNTHETIC[8]
    for line in SYNTHETIC[9:]:
        document.add_paragraph(line)
    document.save(str(path))


def create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Лист1"
    sheet.append(list(HEADERS))
    for row in range(2, 5):
        for col in range(1, 24):
            sheet.cell(row, col, "DEMONSTRATION" if row < 4 else "INSTRUCTIONS")
            sheet.cell(row, col).fill = PatternFill("solid", fgColor="DDEEFF")
    sheet["A1"].font = Font(bold=True)
    sheet.column_dimensions["A"].width = 21
    workbook.save(path)
    workbook.close()


def verify_output(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Лист1"]
        assert sheet.max_row == 4 and sheet.max_column == 23
        assert [sheet.cell(r, 20).value for r in range(2, 5)] == [1, 3, 1]
        assert sheet["N2"].value == 100 and sheet["P2"].value == 100.5
        assert sheet["N4"].value == 103 and sheet["P4"].value == 103.25
        assert sheet["L2"].value == '=IFERROR(K2/J2,"")'
        assert sheet["L3"].value is None
        assert sheet["L4"].value == '=IFERROR(K4/J4,"")'
        assert sheet["M2"].value == "=O2-N2" and sheet["M4"].value == "=O4-N4"
        assert sheet["W3"].data_type == "s" and str(sheet["W3"].value).startswith("=1+1")
        for row in range(2, 5):
            text = str(sheet.cell(row, 23).value)
            assert "от верха." in text and "Всего" not in text and "Геолог" not in text
        assert "0,19–0,23м" in str(sheet["W2"].value)
        assert "Запах\nнефти отсутствует" in str(sheet["W2"].value)
        assert sheet["A1"].font.bold and sheet.column_dimensions["A"].width == 21
        assert sheet["A2"].fill.fgColor.rgb == "00DDEEFF"
        assert sheet["N2"].data_type == "n" and sheet["N2"].number_format == "0.00"
        assert sheet["L2"].number_format == "0.00%"
        assert not any(
            cell.value in {"DEMONSTRATION", "INSTRUCTIONS"} for row in sheet for cell in row
        )
    finally:
        workbook.close()


def frozen_smoke(folder: Path) -> int:
    """Run real binary DOC + DOCX conversion; write a safe external status file."""
    folder.mkdir(parents=True, exist_ok=True)
    report = folder / "smoke-result.json"
    try:
        base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2]))
        fixture = base / "fixtures" / "synthetic.doc"
        if not getattr(sys, "frozen", False):
            fixture = base / "tests" / "fixtures" / "synthetic.doc"
        inputs = folder / "Word с пробелами"
        inputs.mkdir()
        shutil.copyfile(fixture, inputs / "Бинарный пример.doc")
        create_docx(inputs / "Открытый пример.docx")
        template = folder / "Шаблон.xlsx"
        create_template(template)
        results = convert_batch(inputs, template)
        assert len(results) == 2 and all(r.status == "success" for r in results)
        for result in results:
            assert result.output is not None
            verify_output(result.output)
        from .cli import main as cli_main

        cli_output = folder / "cli-output"
        assert (
            cli_main(
                [
                    "--input-dir",
                    str(inputs),
                    "--template",
                    str(template),
                    "--output-dir",
                    str(cli_output),
                    "--json",
                ]
            )
            == 0
        )
        for path in cli_output.glob("*.xlsx"):
            verify_output(path)
        assert len(list(cli_output.glob("*.xlsx"))) == 2
        report.write_text(
            json.dumps(
                {
                    "status": "PASS",
                    "version": __version__,
                    "documents": 2,
                    "layers_each": 3,
                    "backend": "ms-doc-piece-table+olefile",
                    "frozen": bool(getattr(sys, "frozen", False)),
                }
            ),
            encoding="utf-8",
        )
        return 0
    except Exception as exc:
        report.write_text(
            json.dumps({"status": "FAIL", "error_type": type(exc).__name__}), encoding="utf-8"
        )
        return 1
