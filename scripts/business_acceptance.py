"""Private real-input acceptance; independent OOXML text, source hashes and XLSX audit.

For binary DOC supply DOCX copies independently converted by LibreOffice in
--independent-dir. LibreOffice is a verification tool, never an app dependency.
This script writes sensitive material ONLY into the explicitly chosen output.
"""

import argparse
import hashlib
import json
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import asdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from coredescriptionw2e.core_parts import safe_key
from coredescriptionw2e.excel.writer import schema
from coredescriptionw2e.service import batch_summary, convert_batch, input_census, result_summary
from scripts.local_validation import audit

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def canonical(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", text))


def xml_text(element):
    return "".join(
        n.text or ""
        if n.tag == W + "t"
        else "\n"
        if n.tag in {W + "br", W + "cr"}
        else "\t"
        if n.tag == W + "tab"
        else ""
        for n in element.iter()
    )


def independent_descriptions(path, mode):
    """Separate implementation: no application reader, parser or domain model."""
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    if mode == "CORE_INTERVAL":
        records = []
        for row in root.iter(W + "tr"):
            cells = row.findall(W + "tc")
            if len(cells) != 2:
                continue
            left = [
                line.strip() for p in cells[0].iter(W + "p") for line in xml_text(p).splitlines()
            ]
            if any(re.fullmatch(r"Керн\s*(?:№\s*)?\d+", p, re.I) for p in left):
                body = "\n".join(xml_text(p) for p in cells[1].iter(W + "p") if xml_text(p).strip())
                if body and not re.fullmatch(r"Керн\s+не\s+поднят[.!]?", body, re.I):
                    records.append(body)
        return records
    parents = {child: parent for parent in root.iter() for child in parent}
    records = []
    current = None
    table_body = False
    for paragraph, text in (
        (p, line.strip()) for p in root.iter(W + "p") for line in xml_text(p).splitlines()
    ):
        if not text:
            continue
        cell = paragraph
        while cell in parents and cell.tag != W + "tc":
            cell = parents[cell]
        is_last = True
        if cell.tag == W + "tc":
            cells = parents[cell].findall(W + "tc")
            is_last = cells[-1] is cell
        if re.match(r"^(?:Слой\s*№?\s*\d+\b|Интервал\s*№?\s*\d+\.\s*\d)", text, re.I):
            if current is not None:
                records.append("\n".join(current))
            current = []
            table_body = cell.tag == W + "tc" and is_last
        elif re.match(
            r"^(?:Долбление\s*№?\s*\d+|Всего\b|Описание\s+керна\s+(?:составил|выполнил)|Геолог\b|Привязк\w*\s+керна|Начальник\b|Зав\.|Подпись\b|Отбор\s+керна\s+произв)",
            text,
            re.I,
        ):
            if current is not None:
                records.append("\n".join(current))
                current = None
        elif re.match(r"^(?:Пласт\s+\S+|Доюрский\s+комплекс)", text, re.I):
            if current is not None:
                if re.search(r"от\s+верха\s*\.?$", " ".join(current), re.I):
                    records.append("\n".join(current))
                    current = None
                else:
                    current.append(text)
        elif current is not None and (not table_body or is_last):
            current.append(text)
    if current is not None:
        records.append("\n".join(current))
    return records


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--independent-dir", type=Path)
    args = parser.parse_args()
    if args.output_dir.exists():
        parser.error("Use a new unique output directory.")
    census = input_census(args.input_dir)
    hashes = {p.name: digest(p) for p, reason in census if not reason}
    template_hash = digest(args.template)
    results = convert_batch(args.input_dir, args.template, args.output_dir)
    details = []
    for result in results:
        detail = result_summary(result)
        parsed = result.parsed
        detail["input_unchanged"] = hashes.get(result.source.name) == digest(result.source)
        if not parsed or not result.output:
            detail["coverage_status"] = "NOT_RUN"
            detail["blocker"] = [i.code for i in result.issues if i.severity == "error"]
            details.append(detail)
            continue
        workbook = load_workbook(result.output)
        main_sheet = next(
            s
            for s in workbook
            if s.max_column == 24 and s.title not in {"Контроль", "Реестр", "Источники", "Описание"}
        )
        detail["checks"] = audit(
            parsed,
            workbook,
            main_sheet,
            schema(main_sheet, output=True),
            result.source,
            result.output,
        )
        reference = (
            result.source
            if result.source.suffix.lower() == ".docx"
            else args.independent_dir / (result.source.stem + ".docx")
            if args.independent_dir
            else None
        )
        originals = (
            [c.full_text for c in parsed.cores]
            if parsed.cores
            else [layer.full_text for layer in parsed.layers]
        )
        independent = (
            independent_descriptions(reference, parsed.mode)
            if reference and reference.exists()
            else None
        )
        detail["independent_coverage"] = {
            "status": "NOT_RUN"
            if independent is None
            else "PASS"
            if [canonical(t) for t in independent] == [canonical(t) for t in originals]
            else "FAIL",
            "source_descriptions": len(independent) if independent is not None else None,
            "model_descriptions": len(originals),
            "independent_descriptions": independent,
        }
        originals_blocks = (
            [b for c in parsed.cores for b in c.body_blocks]
            if parsed.cores
            else [b for layer in parsed.layers for b in layer.body_blocks]
        )
        parts = [p for c in parsed.active_cores for p in c.parts]
        w = [
            main_sheet.cell(row, schema(main_sheet, output=True)[22]).value
            for row in range(2, main_sheet.max_row + 1)
        ]
        covered = all("\n".join(p.full_text for p in c.parts) == c.full_text for c in parsed.cores)
        duplicates_valid = all(
            safe_key(next(c.full_text for c in parsed.cores if c.source == d.source))
            == safe_key(next(c.full_text for c in parsed.cores if c.source == d.retained_source))
            for d in parsed.suppressed_duplicates
        )
        detail["coverage"] = {
            "parents": len(originals),
            "exported_parts": len(parts),
            "source_blocks": len(originals_blocks),
            "source_numeric_tokens": sum(
                len(re.findall(r"\d+(?:[.,]\d+)?", b.text)) for b in originals_blocks
            ),
            "exported_numeric_tokens": sum(
                len(re.findall(r"\d+(?:[.,]\d+)?", text or "")) for text in w
            ),
            "registered_duplicate_numeric_tokens": sum(
                len(re.findall(r"\d+(?:[.,]\d+)?", c.full_text))
                for c in parsed.cores
                if c.source in {d.source for d in parsed.suppressed_duplicates}
            ),
            "split_complete": covered,
            "duplicates_reproducible": duplicates_valid,
        }
        detail["coverage"]["numeric_balance"] = (
            detail["coverage"]["source_numeric_tokens"]
            == detail["coverage"]["exported_numeric_tokens"]
            + detail["coverage"]["registered_duplicate_numeric_tokens"]
        )
        detail["coverage_status"] = (
            "PASS"
            if all(detail["checks"].values())
            and covered
            and duplicates_valid
            and detail["coverage"]["numeric_balance"]
            and detail["independent_coverage"]["status"] == "PASS"
            else "FAIL"
        )
        detail["rows"] = main_sheet.max_row - 1
        detail["reconciliation"] = parsed.reconciliation
        detail["suppressed_duplicates"] = [asdict(d) for d in parsed.suppressed_duplicates]
        detail["part_map"] = [asdict(p) for p in parts]
        detail["issues_by_code"] = dict(Counter(i.code for i in parsed.issues))
        workbook.close()
        details.append(detail)
    report = {
        "summary": batch_summary(results),
        "template_unchanged": template_hash == digest(args.template),
        "source_hashes": hashes,
        "census": [
            {"file": p.name, "type": p.suffix.casefold(), "reason": reason or "queued"}
            for p, reason in census
        ],
        "documents": details,
        "NOT_RUN": ["Microsoft Excel recalculation", "physical firewall offline smoke"],
        "PENDING": ["Human geological and visual acceptance"],
        "formula_note": "openpyxl does not calculate Excel formulas.",
    }
    (args.output_dir / "report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    lines = [
        "# Приватная приёмка бизнес-логики, партия 1",
        "",
        "Не публиковать: содержит имена документов и локальные результаты.",
        "",
        json.dumps(report["summary"], ensure_ascii=False),
        "",
        "| Документ | Статус | Родители | Части | Строки | Покрытие |",
        "|---|---|---:|---:|---:|---|",
    ]
    for detail in details:
        lines.append(
            f"| {detail['file']} | {detail['document_status']} | {detail.get('core_records') or detail.get('runs')} | {detail.get('core_parts', 0)} | {detail.get('rows', 0)} | {detail['coverage_status']} |"
        )
        if detail.get("output"):
            lines.append(f"\nРезультат: `{Path(str(detail['output'])).resolve()}`\n")
    lines.extend(
        [
            "",
            "Полные проблемы, SourceRef, исходные блоки, карты частей и дубли — в batch JSON, report.json и листах Контроль / Реестр / Источники.",
            "",
            "Итоги J/K — по уникальным родителям. Неизвестная J нулевого интервала не заменяется его длиной. Итоги Word требуют проверки семантики (коробки/линейный вынос).",
            "",
            "NOT_RUN: Microsoft Excel recalculation; physical firewall offline smoke. PENDING: визуальная/геологическая приёмка человеком. Независимый OOXML-аудит не вычисляет формулы.",
        ]
    )
    (args.output_dir / "REPORT.md").write_text("\n".join(lines), encoding="utf-8")
    print(
        json.dumps(
            {
                "summary": report["summary"],
                "coverage": [d["coverage_status"] for d in details],
                "template_unchanged": report["template_unchanged"],
            }
        )
    )
    return int(
        any(d["coverage_status"] == "FAIL" for d in details) or not report["template_unchanged"]
    )


if __name__ == "__main__":
    raise SystemExit(main())
