"""Private local audit. Source text and reports must stay in the ignored data folder."""

import argparse
import json
import os
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from coredescriptionw2e.domain import CorePart, Layer
from coredescriptionw2e.excel.writer import schema
from coredescriptionw2e.export_rows import export_rows
from coredescriptionw2e.parser import parse_document
from coredescriptionw2e.word.common import read_word


def audit(parsed, workbook, sheet, mapping, word, xlsx):
    entries = export_rows(parsed)
    checks = {
        "rows": sheet.max_row == len(entries) + 1,
        "columns": sheet.max_column == 24,
        "control": "Контроль" in workbook.sheetnames,
        "status_ru": workbook["Контроль"]["A1"].value
        in {"Проверено автоматически", "Требует проверки"},
        "review_visible": parsed.status != "REVIEW_REQUIRED" or workbook.active.title == "Контроль",
        "w_complete": True,
        "depth_values": True,
        "recovery_all_run_rows": True,
        "no_invented_layers": True,
        "numbers": True,
        "unknown_metadata_blank": True,
        "source_reference": True,
        "word_well_preserved": True,
        "text_and_numeric_tokens": True,
        "split_coverage": True,
        "internal_links": True,
    }
    for row, (owner, entry) in enumerate(entries, 2):
        original = entry.full_text if entry else None
        w = sheet.cell(row, mapping[22])
        checks["w_complete"] &= w.value == original
        checks["text_and_numeric_tokens"] &= re.findall(
            r"\d+(?:[.,]\d+)?", w.value or ""
        ) == re.findall(r"\d+(?:[.,]\d+)?", original or "")
        for key, value in ((9, owner.drilled), (10, owner.recovered)):
            checks["recovery_all_run_rows"] &= sheet.cell(row, mapping[key]).value == (
                float(value) if value is not None else None
            )
        j, k = (get_column_letter(mapping[key]) for key in (9, 10))
        expected_formula = (
            f'=IFERROR({k}{row}/{j}{row},"")'
            if owner.drilled is not None and owner.drilled > 0 and owner.recovered is not None
            else None
        )
        checks["recovery_all_run_rows"] &= sheet.cell(row, mapping[11]).value == expected_formula
        checks["recovery_all_run_rows"] &= sheet.cell(row, mapping[11]).number_format == "0.00%"
        for key, interval in (
            (7, owner.drilling),
            (13, entry.drilling if isinstance(entry, Layer) else None),
            (15, entry.gis if isinstance(entry, Layer) else None),
        ):
            for offset, value in enumerate(
                (interval.top, interval.bottom) if interval else (None, None)
            ):
                checks["depth_values"] &= sheet.cell(row, mapping[key + offset]).value == (
                    float(value) if value is not None else None
                )
        checks["numbers"] &= sheet.cell(row, mapping[19]).value == (
            entry.number if isinstance(entry, Layer) else None
        )
        if isinstance(entry, CorePart) or entry is None:
            checks["no_invented_layers"] &= all(
                sheet.cell(row, mapping[key]).value is None for key in (12, 13, 14, 15, 16, 19)
            )
        checks["unknown_metadata_blank"] &= all(
            sheet.cell(row, mapping[key]).value is None for key in (1, 2, 4, 5, 17)
        )
        checks["source_reference"] &= sheet.cell(row, mapping[6]).value == os.path.relpath(
            word.resolve(), xlsx.parent.resolve()
        )
        checks["word_well_preserved"] &= (
            sheet.cell(row, mapping[3]).value == parsed.metadata.well == sheet.cell(row, 24).value
        )
    for core in parsed.cores:
        checks["split_coverage"] &= "\n".join(p.full_text for p in core.parts) == core.full_text
        checks["split_coverage"] &= [
            b for p in core.parts for b in p.body_blocks
        ] == core.body_blocks
    for ws in workbook:
        for cells in ws:
            for cell in cells:
                if cell.hyperlink:
                    target = cell.hyperlink.target
                    if not target or not target.startswith("#'"):
                        checks["internal_links"] = False
                        continue
                    title, address = target[2:].rsplit("'!", 1)
                    title = title.replace("''", "'")
                    checks["internal_links"] &= title in workbook.sheetnames
                    if title in workbook.sheetnames:
                        checks["internal_links"] &= workbook[title][address].value is not None
    return checks


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--word", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--expected", type=Path)
    parser.add_argument("--reference-docx", type=Path)
    parser.add_argument("--template", type=Path)
    args = parser.parse_args()
    parsed = parse_document(read_word(args.word))
    workbook = load_workbook(args.xlsx, data_only=False)
    try:
        sheet = next(
            s
            for s in workbook
            if s.max_column == 24 and s.title not in {"Контроль", "Описание", "Реестр", "Источники"}
        )
        mapping = schema(sheet, output=True)
        checks = audit(parsed, workbook, sheet, mapping, args.word, args.xlsx)
        if args.template:
            template = load_workbook(args.template)
            checks["header_preserved"] = [c.value for c in sheet[1]][:23] == [
                c.value for c in template.active[1]
            ]
            template.close()
        if args.reference_docx:
            reference = read_word(args.reference_docx)
            checks["independent_text_match"] = re.sub(
                r"\s+", "", parsed.document.raw_text
            ) == re.sub(r"\s+", "", reference.raw_text)
        summary = {
            "mode": parsed.mode,
            "runs": len(parsed.runs),
            "layers": len(parsed.layers),
            "core_records": len(parsed.cores),
            "core_parts": sum(len(c.parts) for c in parsed.active_cores),
            "no_core": len(parsed.no_core),
            "suppressed_duplicates": len(parsed.suppressed_duplicates),
            "distribution": [len(r.layers) for r in parsed.runs],
            "issues": dict(Counter(i.code for i in parsed.issues)),
            "reconciliation": parsed.reconciliation,
        }
        if args.expected:
            expected = json.loads(args.expected.read_text(encoding="utf-8"))
            checks["expected_structure"] = all(summary.get(k) == v for k, v in expected.items())
        summary["checks"] = checks
        summary["status"] = parsed.status if all(checks.values()) and parsed.exportable else "FAIL"
        args.xlsx.with_suffix(".validation.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        print(
            json.dumps(
                {
                    k: summary[k]
                    for k in ("mode", "status", "runs", "layers", "core_records", "checks")
                },
                ensure_ascii=False,
            )
        )
        return 0 if summary["status"] in {"PASS", "REVIEW_REQUIRED"} else 1
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
